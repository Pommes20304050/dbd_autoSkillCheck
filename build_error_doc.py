"""Assemble the per-chapter Markdown error reference into a polished Word doc.

Reads:
    error_doc_chapters/01_server.md  ...  10_setup_runtime.md

Writes:
    DBD_Auto_Skill_Check_Error_Reference.docx

Design goals:
- Cover page with large title and accent bar
- Master Table of Contents with hyperlinks: lists every chapter AND every
  error code with its title (clickable, jumps directly to the entry)
- Each chapter starts on a new page with a numbered chapter header and
  a per-chapter mini-TOC of all error codes in that chapter
- Each error entry has a bookmark, an accent code badge, and a clean layout
- Footer with centered page numbers
- Code blocks rendered in Consolas with light grey shading

Run:
    python build_error_doc.py
"""
from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------
CHAPTERS_DIR = Path(__file__).parent / "error_doc_chapters"
OUT_PATH = Path(__file__).parent / "DBD_Auto_Skill_Check_Error_Reference.docx"

CHAPTER_FILES: list[tuple[str, str]] = [
    ("01_server.md",            "Server"),
    ("02_inference_worker.md",  "Inference Worker"),
    ("03_ai_model.md",          "AI Model"),
    ("04_screen_capture.md",    "Screen Capture"),
    ("05_system_info.md",       "System Information"),
    ("06_perf_monitor.md",      "Performance Monitor"),
    ("07_fps_advisor.md",       "FPS Advisor"),
    ("08_env_preflight.md",     "Environment & Preflight"),
    ("09_frontend.md",          "Frontend"),
    ("10_setup_runtime.md",     "Setup & Runtime"),
]

# Orange/black accent palette (matches the project's UI).
ACCENT       = "FF6B2C"
ACCENT_DARK  = "CC551F"
TEXT_DARK    = "1A1A1A"
TEXT_MUTED   = "606060"
CODE_BG      = "F4F4F4"
DIVIDER_BG   = "FFE6D7"


# ---------------------------------------------------------------------------
# XML helpers (python-docx exposes only a subset of Word features)
# ---------------------------------------------------------------------------
def _set(el, **attrs):
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), v)
    return el


def _new(tag: str, **attrs):
    el = OxmlElement(f"w:{tag}")
    return _set(el, **attrs)


def shade(paragraph, hex_color: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = _new("shd", val="clear", color="auto", fill=hex_color)
    pPr.append(shd)


def border(paragraph, *, side: str, color: str, size: int = 8, space: int = 4) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = _new("pBdr")
        pPr.append(pBdr)
    el = _new(side, val="single", sz=str(size), space=str(space), color=color)
    pBdr.append(el)


def add_bookmark(paragraph, name: str, bk_id: int) -> None:
    bm_start = _new("bookmarkStart", id=str(bk_id), name=name)
    bm_end = _new("bookmarkEnd", id=str(bk_id))
    paragraph._p.insert(0, bm_start)
    paragraph._p.append(bm_end)


def add_hyperlink(paragraph, anchor: str, text: str, *,
                  bold: bool = False, color: str | None = None,
                  size_pt: int | None = None, font: str | None = None,
                  underline: bool = False) -> None:
    hl = _new("hyperlink", anchor=anchor)
    r = _new("r")
    rPr = _new("rPr")
    if bold:
        rPr.append(_new("b"))
    if underline:
        rPr.append(_new("u", val="single"))
    if color is not None:
        rPr.append(_new("color", val=color))
    if size_pt is not None:
        rPr.append(_new("sz", val=str(size_pt * 2)))
    if font is not None:
        rPr.append(_new("rFonts", ascii=font, hAnsi=font, cs=font))
    r.append(rPr)
    t = _new("t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    hl.append(r)
    paragraph._p.append(hl)


def add_tab_with_dotted_leader(paragraph, right_pos_inches: float) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    tabs = _new("tabs")
    tabs.append(_new("tab", val="right", leader="dot",
                     pos=str(int(right_pos_inches * 1440))))
    pPr.append(tabs)


def add_page_number_field_to_footer(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ""
    run = p.add_run()
    run._r.append(_new("fldChar", fldCharType="begin"))
    instr = _new("instrText"); instr.text = " PAGE "
    run._r.append(instr)
    run._r.append(_new("fldChar", fldCharType="separate"))
    t = _new("t"); t.text = "1"
    run._r.append(t)
    run._r.append(_new("fldChar", fldCharType="end"))
    for r in p.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(TEXT_MUTED)


def suppress_update_fields_prompt(doc: Document) -> None:
    """Tell Word NOT to ask 'update fields from external sources?' on open.
    Page numbers (PAGE fields) still self-update on display; the manual
    hyperlinked TOC doesn't need updating at all."""
    settings = doc.settings.element
    el = settings.find(qn("w:updateFields"))
    if el is None:
        el = _new("updateFields", val="false")
        settings.append(el)
    else:
        el.set(qn("w:val"), "false")


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------
def setup_styles(doc: Document) -> None:
    styles = doc.styles

    base = styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(11)
    base.font.color.rgb = RGBColor.from_string(TEXT_DARK)
    base.paragraph_format.space_after = Pt(4)
    base.paragraph_format.line_spacing = 1.25

    for level, size, color in [
        ("Heading 1", 24, ACCENT_DARK),
        ("Heading 2", 16, ACCENT_DARK),
        ("Heading 3", 13, TEXT_DARK),
        ("Heading 4", 11, TEXT_DARK),
    ]:
        s = styles[level]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after = Pt(6)

    if "ErrorCode" not in [s.name for s in styles]:
        s = styles.add_style("ErrorCode", WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = styles["Normal"]
        s.font.name = "Consolas"
        s.font.size = Pt(9)
        s.font.color.rgb = RGBColor.from_string(TEXT_DARK)
        s.paragraph_format.space_before = Pt(0)
        s.paragraph_format.space_after = Pt(0)
        s.paragraph_format.left_indent = Inches(0.25)
        s.paragraph_format.line_spacing = 1.15

    if "TOCEntry" not in [s.name for s in styles]:
        s = styles.add_style("TOCEntry", WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = styles["Normal"]
        s.font.size = Pt(10)
        s.paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------------------
# Markdown parsing — line-by-line, tuned to what the chapter agents emit.
# ---------------------------------------------------------------------------
INLINE_RE   = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
ERROR_HEAD  = re.compile(r"^##\s+(?P<code>E\.[A-Z]+\.\d+[A-Za-z]?)\s*[—\-:]\s*(?P<title>.+)$")
CHAPTER_HEAD = re.compile(r"^#\s+(?:Chapter\s+\d+\s*[—\-:]\s*)?(?P<title>.+)$")


def _add_runs(paragraph, text: str) -> None:
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            r = paragraph.add_run(chunk[2:-2])
            r.bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            r = paragraph.add_run(chunk[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            r = paragraph.add_run(chunk[1:-1])
            r.italic = True
        else:
            paragraph.add_run(chunk)


def _scan_chapter(md_text: str) -> tuple[str, list[tuple[str, str]]]:
    """Return (chapter_title, [(code, title), ...]) without heavy parsing."""
    title = None
    entries: list[tuple[str, str]] = []
    for raw in md_text.splitlines():
        line = raw.strip()
        if title is None:
            m = CHAPTER_HEAD.match(line)
            if m:
                title = m.group("title").strip()
                continue
        m = ERROR_HEAD.match(line)
        if m:
            entries.append((m.group("code").strip(), m.group("title").strip()))
    return title or "Untitled Chapter", entries


def _slug(code: str) -> str:
    """Word bookmark names: letters, digits, underscore. No dots or dashes."""
    return "err_" + re.sub(r"[^A-Za-z0-9]+", "_", code)


# ---------------------------------------------------------------------------
# Cover page
# ---------------------------------------------------------------------------
def emit_cover(doc: Document, total_errors: int) -> None:
    # Top spacer
    doc.add_paragraph("").paragraph_format.space_after = Pt(80)

    # Accent bar above title
    bar = doc.add_paragraph()
    bar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(bar, ACCENT)
    bar.paragraph_format.space_after = Pt(18)
    r = bar.add_run("   ")
    r.font.size = Pt(2)

    # Big title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DBD Auto Skill Check")
    r.bold = True
    r.font.size = Pt(36)
    r.font.color.rgb = RGBColor.from_string(TEXT_DARK)

    # Subtitle
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Complete Error & Troubleshooting Reference")
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
    p.paragraph_format.space_after = Pt(36)

    # Italic strapline
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Every error this tool can raise — what it means, why it "
                  "happened, and exactly how to fix it.")
    r.italic = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(TEXT_MUTED)
    p.paragraph_format.space_after = Pt(60)

    # Stats box
    box = doc.add_paragraph()
    box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(box, DIVIDER_BG)
    box.paragraph_format.space_before = Pt(6)
    box.paragraph_format.space_after = Pt(6)
    r = box.add_run(f"  {len(CHAPTER_FILES)} chapters    "
                    f"{total_errors} catalogued errors    "
                    f"Generated {date.today().isoformat()}  ")
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(ACCENT_DARK)

    # Bottom accent bar
    p = doc.add_paragraph("")
    p.paragraph_format.space_before = Pt(80)
    bar = doc.add_paragraph()
    bar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(bar, ACCENT)
    r = bar.add_run("   ")
    r.font.size = Pt(2)

    doc.add_page_break()


# ---------------------------------------------------------------------------
# Master Table of Contents
# ---------------------------------------------------------------------------
def emit_master_toc(doc: Document, chapters: list[dict]) -> None:
    h = doc.add_paragraph()
    r = h.add_run("Table of Contents")
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
    border(h, side="bottom", color=ACCENT, size=12, space=6)
    h.paragraph_format.space_after = Pt(18)

    intro = doc.add_paragraph()
    r = intro.add_run("Click any entry to jump directly to that section. "
                      "Each chapter also has its own mini-table of contents "
                      "listing every error in that chapter.")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(TEXT_MUTED)
    intro.paragraph_format.space_after = Pt(18)

    for ch in chapters:
        # Chapter line — bold orange link with dotted leader
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        add_tab_with_dotted_leader(p, right_pos_inches=6.3)
        add_hyperlink(
            p,
            anchor=f"chapter_{ch['num']:02d}",
            text=f"Chapter {ch['num']} — {ch['title']}",
            bold=True, color=ACCENT_DARK, size_pt=13,
        )
        # Tab + error count on the right
        tab_run = p.add_run("\t")
        tab_run.font.size = Pt(11)
        cnt = p.add_run(f"{len(ch['entries'])} entries")
        cnt.font.size = Pt(10)
        cnt.italic = True
        cnt.font.color.rgb = RGBColor.from_string(TEXT_MUTED)

        # Error code lines
        for code, title in ch["entries"]:
            p = doc.add_paragraph(style="TOCEntry")
            p.paragraph_format.left_indent = Inches(0.4)
            add_tab_with_dotted_leader(p, right_pos_inches=6.3)
            # Error code in mono accent
            add_hyperlink(p, anchor=_slug(code), text=code,
                          color=ACCENT_DARK, size_pt=10, font="Consolas")
            # Title in normal weight
            p.add_run("  ").font.size = Pt(10)
            add_hyperlink(p, anchor=_slug(code), text=title,
                          color=TEXT_DARK, size_pt=10)

    doc.add_page_break()


# ---------------------------------------------------------------------------
# Chapter intro page (mini-TOC)
# ---------------------------------------------------------------------------
def emit_chapter_intro(doc: Document, ch: dict) -> None:
    # Bookmarked anchor for the master TOC
    p = doc.add_paragraph()
    add_bookmark(p, name=f"chapter_{ch['num']:02d}", bk_id=10000 + ch["num"])
    r = p.add_run(f"CHAPTER {ch['num']:02d}")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(ACCENT)
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    r = p.add_run(ch["title"])
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor.from_string(TEXT_DARK)
    border(p, side="bottom", color=ACCENT, size=14, space=8)
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    r = p.add_run(f"{len(ch['entries'])} catalogued errors in this chapter.")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(TEXT_MUTED)
    p.paragraph_format.space_after = Pt(12)

    p = doc.add_paragraph()
    r = p.add_run("In this chapter")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
    p.paragraph_format.space_after = Pt(6)

    for code, title in ch["entries"]:
        p = doc.add_paragraph(style="TOCEntry")
        add_tab_with_dotted_leader(p, right_pos_inches=6.3)
        add_hyperlink(p, anchor=_slug(code), text=code,
                      color=ACCENT_DARK, size_pt=10, font="Consolas")
        p.add_run("  ").font.size = Pt(10)
        add_hyperlink(p, anchor=_slug(code), text=title,
                      color=TEXT_DARK, size_pt=10)

    doc.add_page_break()


# ---------------------------------------------------------------------------
# Render markdown body of a chapter — skips the chapter H1 (already drawn).
# ---------------------------------------------------------------------------
def emit_chapter_body(doc: Document, md_text: str, bk_id_start: int) -> int:
    bk_id = bk_id_start
    lines = md_text.splitlines()
    in_code = False
    code_buf: list[str] = []
    saw_chapter_heading = False

    def flush_code():
        if not code_buf:
            return
        for ln in code_buf:
            p = doc.add_paragraph(style="ErrorCode")
            p.add_run(ln if ln else " ")
            shade(p, CODE_BG)
            border(p, side="left", color=ACCENT, size=12, space=6)
        code_buf.clear()

    for raw in lines:
        line = raw.rstrip()

        # Code fence toggling
        if line.strip().startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_buf.append(line)
            continue

        if not line.strip():
            doc.add_paragraph("").paragraph_format.space_after = Pt(2)
            continue

        # Headings
        if line.startswith("# "):
            # Skip the chapter heading — we already drew the chapter intro.
            if not saw_chapter_heading:
                saw_chapter_heading = True
                continue
            # Defensive: if a later H1 shows up, still render it.
            doc.add_heading(line[2:].strip(), level=1)
            continue

        if line.startswith("## "):
            stripped = line[3:].strip()
            m = ERROR_HEAD.match(line)
            if m:
                code = m.group("code").strip()
                title = m.group("title").strip()
                # Code badge
                badge = doc.add_paragraph()
                badge.paragraph_format.space_before = Pt(18)
                badge.paragraph_format.space_after = Pt(0)
                add_bookmark(badge, name=_slug(code), bk_id=bk_id)
                bk_id += 1
                r = badge.add_run(code)
                r.bold = True
                r.font.name = "Consolas"
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor.from_string(ACCENT)
                # Title
                p = doc.add_paragraph()
                r = p.add_run(title)
                r.bold = True
                r.font.size = Pt(15)
                r.font.color.rgb = RGBColor.from_string(TEXT_DARK)
                border(p, side="bottom", color=ACCENT, size=6, space=4)
                p.paragraph_format.space_after = Pt(6)
            else:
                doc.add_heading(stripped, level=2)
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=4)
            continue

        # Lists
        if line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, line[2:])
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, re.sub(r"^\d+\.\s", "", line))
            continue

        # Blockquote
        if line.startswith("> "):
            p = doc.add_paragraph()
            shade(p, DIVIDER_BG)
            border(p, side="left", color=ACCENT, size=18, space=6)
            r = p.add_run(line[2:])
            r.italic = True
            continue

        # Horizontal rule
        if line.strip() in ("---", "***"):
            sep = doc.add_paragraph()
            border(sep, side="bottom", color=ACCENT, size=6, space=2)
            continue

        # Plain text
        p = doc.add_paragraph()
        _add_runs(p, line)

    if in_code:
        flush_code()

    doc.add_page_break()
    return bk_id


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> int:
    if not CHAPTERS_DIR.is_dir():
        print(f"[ERROR] Chapters directory missing: {CHAPTERS_DIR}", file=sys.stderr)
        return 1

    chapters: list[dict] = []
    for idx, (fname, default_title) in enumerate(CHAPTER_FILES, start=1):
        path = CHAPTERS_DIR / fname
        if not path.exists():
            print(f"[WARN] Missing chapter file: {fname}")
            chapters.append({
                "num": idx, "title": default_title, "path": None,
                "entries": [], "missing": True,
            })
            continue
        text = path.read_text(encoding="utf-8")
        scan_title, entries = _scan_chapter(text)
        title = default_title or scan_title
        chapters.append({
            "num": idx, "title": title, "path": path,
            "entries": entries, "missing": False, "text": text,
        })

    total_errors = sum(len(c["entries"]) for c in chapters)
    print(f"[INFO] {len(chapters)} chapters, {total_errors} error entries")

    doc = Document()
    setup_styles(doc)

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        add_page_number_field_to_footer(section)

    emit_cover(doc, total_errors)
    emit_master_toc(doc, chapters)

    bk_id = 100000
    for ch in chapters:
        emit_chapter_intro(doc, ch)
        if ch["missing"]:
            p = doc.add_paragraph()
            r = p.add_run("Chapter file not found — re-run chapter generation.")
            r.italic = True
            r.font.color.rgb = RGBColor.from_string(TEXT_MUTED)
            doc.add_page_break()
            continue
        bk_id = emit_chapter_body(doc, ch["text"], bk_id)

    suppress_update_fields_prompt(doc)
    doc.save(str(OUT_PATH))
    print(f"[OK] Wrote {OUT_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
